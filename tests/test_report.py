import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from typer.testing import CliRunner

from data_migration_concept.checks import run_checks
from data_migration_concept.cli import app
from data_migration_concept.config import load_project
from data_migration_concept.extract import extract_object
from data_migration_concept.profile import profile_object
from data_migration_concept.recon import run_recons
from data_migration_concept.report import ReportError, collect_report, render_html, write_reports
from data_migration_concept.runlog import RunLog
from data_migration_concept.transform import transform_object

runner = CliRunner()


@pytest.fixture
def working():
    engine = create_engine("sqlite://")
    log = RunLog(engine)
    log.ensure_schema()
    return engine, log


def _edit(project_dir: Path, rel: str, old: str, new: str):
    path = project_dir / "objects" / "supplier" / rel
    content = path.read_text(encoding="utf-8")
    assert old in content, f"{old!r} not in {rel}"
    path.write_text(content.replace(old, new), encoding="utf-8")


def _pipeline(project_dir: Path, engine, log):
    """Each stage gets its own run, as the CLI does."""
    project = load_project(project_dir)
    name = project.config.name
    extract_run = log.start_run(name, "supplier")
    extract_object(project, "supplier", engine, log, extract_run)
    transform_run = log.start_run(name, "supplier")
    transform_object(project, "supplier", engine, log, transform_run, write_build=False)
    check_run = log.start_run(name, "supplier")
    run_checks(project, "supplier", engine, log, check_run)
    recon_run = log.start_run(name, "supplier")
    run_recons(project, "supplier", engine, log, recon_run)
    profile_run = log.start_run(name, "supplier")
    profile_object(project, "supplier", engine, log, profile_run)
    return project, {
        "extract": extract_run,
        "transform": transform_run,
        "checks": check_run,
        "recon": recon_run,
        "profile": profile_run,
    }


def test_report_combines_the_latest_run_of_each_kind(example_project, working):
    engine, log = working
    project, runs = _pipeline(example_project, engine, log)
    data = collect_report(log, project.config.name, "supplier")
    assert data["overall"] == "OK"
    assert {r["run_id"] for r in data["runs"]} == set(runs.values())
    assert {s["kind"] for s in data["steps"]} == {
        "extract",
        "transform",
        "checks",
        "recon",
        "profile",
    }
    s = data["summary"]
    assert s["checks_passed"] == s["checks_total"] > 0
    assert s["recons_passed"] == s["recons_total"] == 4
    assert list(data["profiles"]) == ["stg_supplier_lfa1"]
    assert data["rejects"] == []
    json.dumps(data)  # must be serialisable as is


def test_newer_run_replaces_older_one_of_the_same_kind(example_project, working):
    engine, log = working
    project, runs = _pipeline(example_project, engine, log)
    newer = log.start_run(project.config.name, "supplier")
    run_checks(project, "supplier", engine, log, newer)
    data = collect_report(log, project.config.name, "supplier")
    ids = {r["run_id"] for r in data["runs"]}
    assert newer in ids and runs["checks"] not in ids


def test_blocking_failures_give_blocked_with_capped_rejects(project_copy, working):
    engine, log = working
    data_file = project_copy / "objects" / "supplier" / "data" / "lfa1.csv"
    data_file.write_text(
        data_file.read_text(encoding="utf-8") + "0000100001,Dup A,DE\n0000100002,Dup B,DE\n",
        encoding="utf-8",
    )
    project, _ = _pipeline(project_copy, engine, log)
    data = collect_report(log, project.config.name, "supplier", max_rejects=2)
    assert data["overall"] == "BLOCKED"
    assert data["summary"]["checks_blocking"] == 1
    unique = [r for r in data["rejects"] if r["check_name"] == "partner_unique"]
    assert len(unique) == 2
    failing = next(c for c in data["checks"] if c["check_name"] == "partner_unique")
    assert failing["failed_count"] == 4 and failing["rejects_shown"] == 2


def test_warnings_give_attention(project_copy, working):
    engine, log = working
    _edit(project_copy, "checks.yaml", "WHERE LENGTH", "WHERE 1 = 1 OR LENGTH")
    project, _ = _pipeline(project_copy, engine, log)
    data = collect_report(log, project.config.name, "supplier")
    assert data["overall"] == "ATTENTION" and data["summary"]["checks_warnings"] == 1


def test_report_for_a_single_run(example_project, working):
    engine, log = working
    project, runs = _pipeline(example_project, engine, log)
    data = collect_report(log, project.config.name, "supplier", run_id=runs["recon"])
    assert [r["run_id"] for r in data["runs"]] == [runs["recon"]]
    assert data["summary"]["recons_total"] == 4 and data["summary"]["checks_total"] == 0


def test_unknown_run_and_empty_log(example_project, working):
    engine, log = working
    with pytest.raises(ReportError, match="unknown run"):
        collect_report(log, "p", "supplier", run_id="nope")
    with pytest.raises(ReportError, match="nothing recorded"):
        collect_report(log, "p", "supplier")


def test_excel_report_has_sheets_and_never_stores_formulas(project_copy, working, tmp_path):
    engine, log = working
    _edit(project_copy, "data/lfa1.csv", "0000100003", "=1+1")
    project, _ = _pipeline(project_copy, engine, log)
    data = collect_report(log, project.config.name, "supplier")
    (path,) = write_reports(data, tmp_path / "out", ("xlsx",))
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Summary",
        "Steps",
        "Checks",
        "Reconciliation",
        "Rejects",
        "Profile",
    ]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows()}
    assert summary["Object"] == "supplier" and summary["Overall"] == "BLOCKED"
    rejects = wb["Rejects"]
    values = [(c.value, c.data_type) for row in rejects.iter_rows(min_row=2) for c in row]
    assert ("=1+1", "s") in values  # kept as text, not a formula
    assert not any(c.data_type == "f" for ws in wb for row in ws.iter_rows() for c in row)
    checks = {row[0].value: row[5].value for row in wb["Checks"].iter_rows(min_row=2)}
    assert checks["partner_is_numeric"] == "FAIL"
    profile_cols = [row[1].value for row in wb["Profile"].iter_rows(min_row=2)]
    assert profile_cols == ["LIFNR", "NAME1", "LAND1"]


def test_html_report_escapes_data(project_copy, working, tmp_path):
    engine, log = working
    _edit(project_copy, "data/lfa1.csv", "0000100003", "<b>x</b>")
    project, _ = _pipeline(project_copy, engine, log)
    data = collect_report(log, project.config.name, "supplier")
    html = render_html(data)
    assert "<b>x</b>" not in html and "&lt;b&gt;x&lt;/b&gt;" in html
    assert "<script" not in html
    assert "BLOCKED" in html and "partner_is_numeric" in html
    (path,) = write_reports(data, tmp_path, ("html",))
    assert path.name == "supplier_report.html" and path.read_text(encoding="utf-8") == html


def test_json_report_roundtrips_and_unknown_format(example_project, working, tmp_path):
    engine, log = working
    project, _ = _pipeline(example_project, engine, log)
    data = collect_report(log, project.config.name, "supplier")
    (path,) = write_reports(data, tmp_path, ("json",))
    assert json.loads(path.read_text(encoding="utf-8"))["overall"] == "OK"
    with pytest.raises(ReportError, match="unknown format"):
        write_reports(data, tmp_path, ("pdf",))


def _cli_pipeline(project_dir: Path, monkeypatch, tmp_path: Path):
    monkeypatch.setenv("DMC_WORKING_DB_URL", f"sqlite:///{tmp_path / 'work.db'}")
    for cmd in ("extract", "transform", "check", "load", "recon", "profile"):
        result = runner.invoke(app, [cmd, str(project_dir)])
        assert result.exit_code == 0, (cmd, result.output)


def test_report_command_writes_all_formats(example_project, monkeypatch, tmp_path):
    project_dir = tmp_path / "proj"
    shutil.copytree(example_project, project_dir)
    _cli_pipeline(project_dir, monkeypatch, tmp_path)
    result = runner.invoke(app, ["report", str(project_dir)])
    assert result.exit_code == 0, result.output
    assert "supplier: OK" in result.stdout
    out = project_dir / "reports"
    assert sorted(p.name for p in out.iterdir()) == [
        "supplier_report.html",
        "supplier_report.json",
        "supplier_report.xlsx",
    ]
    data = json.loads((out / "supplier_report.json").read_text(encoding="utf-8"))
    assert {s["kind"] for s in data["steps"]} >= {"extract", "transform", "load", "recon"}


def test_report_command_options(example_project, monkeypatch, tmp_path):
    project_dir = tmp_path / "proj"
    shutil.copytree(example_project, project_dir)
    _cli_pipeline(project_dir, monkeypatch, tmp_path)
    out = tmp_path / "elsewhere"
    result = runner.invoke(
        app, ["report", str(project_dir), "supplier", "--format", "json", "--out", str(out)]
    )
    assert result.exit_code == 0 and [p.name for p in out.iterdir()] == ["supplier_report.json"]

    bad_format = runner.invoke(app, ["report", str(project_dir), "--format", "pdf"])
    assert bad_format.exit_code == 1 and "unknown format" in bad_format.output
    bad_run = runner.invoke(app, ["report", str(project_dir), "--run", "nope"])
    assert bad_run.exit_code == 1 and "unknown run" in bad_run.output


def test_report_command_with_nothing_recorded(example_project, monkeypatch, tmp_path):
    monkeypatch.setenv("DMC_WORKING_DB_URL", f"sqlite:///{tmp_path / 'empty.db'}")
    result = runner.invoke(app, ["report", str(example_project)])
    assert result.exit_code == 1 and "nothing recorded" in result.output
