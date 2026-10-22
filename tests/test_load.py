import csv
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from data_migration_concept.checks import CheckRunError
from data_migration_concept.config import ConfigError, load_project
from data_migration_concept.extract import extract_object
from data_migration_concept.load import LoadBlockedError, LoadError, load_object
from data_migration_concept.runlog import RunLog
from data_migration_concept.transform import transform_object


@pytest.fixture
def working():
    engine = create_engine("sqlite://")
    log = RunLog(engine)
    log.ensure_schema()
    return engine, log


def _edit(project_dir: Path, rel: str, old: str, new: str):
    path = project_dir / "objects" / "supplier" / rel
    text_ = path.read_text(encoding="utf-8")
    assert old in text_, f"{old!r} not in {rel}"
    path.write_text(text_.replace(old, new), encoding="utf-8")


def _prepare(project_dir: Path, engine, log):
    project = load_project(project_dir)
    run_id = log.start_run("p", "supplier")
    extract_object(project, "supplier", engine, log, run_id)
    transform_object(project, "supplier", engine, log, run_id, write_build=False)
    return project


def _load(project, engine, log):
    return load_object(project, "supplier", engine, log, log.start_run("p", "supplier"))


def test_csv_load_writes_clean_file_without_lineage(project_copy, working):
    engine, log = working
    project = _prepare(project_copy, engine, log)
    results = _load(project, engine, log)
    assert [(r.structure, r.kind, r.rows) for r in results] == [("general", "file", 4)]
    out = project_copy / "output" / "supplier_general.csv"
    rows = list(csv.DictReader(out.open(encoding="utf-8")))
    assert len(rows) == 4
    assert list(rows[0]) == [
        "PARTNER",
        "BU_GROUP",
        "NAME_ORG1",
        "COUNTRY",
        "SEARCH_TERM",
        "LEGACY_REF",
    ]
    assert rows[0]["PARTNER"] == "0000100001"  # leading zeros kept
    assert rows[1]["COUNTRY"] == "GB"
    assert not list(out.parent.glob("*.tmp"))


def test_blocking_check_failure_stops_the_load_and_writes_nothing(project_copy, working):
    engine, log = working
    data = project_copy / "objects" / "supplier" / "data" / "lfa1.csv"
    data.write_text(data.read_text(encoding="utf-8") + "0000100001,Dup,DE\n", encoding="utf-8")
    project = _prepare(project_copy, engine, log)
    with pytest.raises(LoadBlockedError, match="partner_unique") as exc:
        _load(project, engine, log)
    assert [r.name for r in exc.value.blocking] == ["partner_unique"]
    assert not (project_copy / "output").exists()


def test_failed_load_keeps_previous_output(project_copy, working):
    engine, log = working
    project = _prepare(project_copy, engine, log)
    _load(project, engine, log)
    out = project_copy / "output" / "supplier_general.csv"
    before = out.read_text(encoding="utf-8")

    data = project_copy / "objects" / "supplier" / "data" / "lfa1.csv"
    data.write_text(data.read_text(encoding="utf-8") + "0000100001,Dup,DE\n", encoding="utf-8")
    project = _prepare(project_copy, engine, log)
    with pytest.raises(LoadBlockedError):
        _load(project, engine, log)
    assert out.read_text(encoding="utf-8") == before


def test_warnings_do_not_block_the_load(project_copy, working):
    engine, log = working
    _edit(project_copy, "checks.yaml", "max_pct: 5", "max_pct: 0")  # still passes: no empties
    _edit(project_copy, "checks.yaml", "WHERE LENGTH", "WHERE 1 = 1 OR LENGTH")  # warn fails
    project = _prepare(project_copy, engine, log)
    assert _load(project, engine, log)[0].rows == 4


def test_load_requires_transform_first(project_copy, working):
    engine, log = working
    project = load_project(project_copy)
    with pytest.raises(CheckRunError, match="run transform first"):
        _load(project, engine, log)


def test_object_without_loads_is_rejected(project_copy, working):
    engine, log = working
    obj = project_copy / "objects" / "supplier" / "object.yaml"
    content = obj.read_text(encoding="utf-8")
    obj.write_text(content[: content.index("loads:")], encoding="utf-8")
    project = _prepare(project_copy, engine, log)
    with pytest.raises(LoadError, match="no 'loads' defined"):
        _load(project, engine, log)


def test_xlsx_and_parquet_and_custom_csv_options(project_copy, working):
    engine, log = working
    _edit(
        project_copy,
        "object.yaml",
        "loads:\n",
        "loads:\n"
        "  - kind: file\n    structure: general\n    path: output/g.xlsx\n    format: xlsx\n"
        "  - kind: file\n    structure: general\n    path: output/g.parquet\n    format: parquet\n"
        "  - kind: file\n    structure: general\n    path: output/g.txt\n    format: csv\n"
        "    delimiter: ';'\n    include_lineage: true\n",
    )
    project = _prepare(project_copy, engine, log)
    assert [r.rows for r in _load(project, engine, log)] == [4, 4, 4, 4]
    out = project_copy / "output"
    ws = load_workbook(out / "g.xlsx").active
    assert [c.value for c in ws[1]][:2] == ["PARTNER", "BU_GROUP"]
    assert ws.max_row == 5
    assert pl.read_parquet(out / "g.parquet").height == 4
    header = (out / "g.txt").read_text(encoding="utf-8").splitlines()[0]
    assert header.startswith("_dmc_row_no;_dmc_run_id;PARTNER")


def test_empty_target_still_writes_header(project_copy, working):
    engine, log = working
    (project_copy / "objects" / "supplier" / "data" / "lfa1.csv").write_text(
        "LIFNR,NAME1,LAND1\n", encoding="utf-8"
    )
    project = _prepare(project_copy, engine, log)
    # no rows: mandatory/unique pass vacuously
    assert _load(project, engine, log)[0].rows == 0
    out = project_copy / "output" / "supplier_general.csv"
    assert out.read_text(encoding="utf-8").splitlines() == [
        "PARTNER,BU_GROUP,NAME_ORG1,COUNTRY,SEARCH_TERM,LEGACY_REF"
    ]


def test_table_load_replace_and_append(project_copy, working, tmp_path, monkeypatch):
    engine, log = working
    dest = tmp_path / "dest.db"
    monkeypatch.setenv("DEST_DB_URL", f"sqlite:///{dest}")
    (project_copy / "project.yaml").write_text(
        (project_copy / "project.yaml").read_text(encoding="utf-8")
        + "  dest:\n    url_env: DEST_DB_URL\n",
        encoding="utf-8",
    )
    _edit(
        project_copy,
        "object.yaml",
        "loads:\n",
        "loads:\n"
        "  - kind: table\n    structure: general\n    connection: dest\n    table: bp_out\n",
    )
    project = _prepare(project_copy, engine, log)
    _load(project, engine, log)
    _load(project, engine, log)  # replace: still 4 rows
    dest_engine = create_engine(f"sqlite:///{dest}")
    with dest_engine.connect() as conn:
        assert conn.execute(text("select count(*) from bp_out")).scalar() == 4
        assert (
            conn.execute(
                text('select "COUNTRY" from bp_out where "PARTNER"=\'0000100002\'')
            ).scalar()
            == "GB"
        )

    _edit(project_copy, "object.yaml", "table: bp_out\n", "table: bp_out\n    mode: append\n")
    project = load_project(project_copy)
    _load(project, engine, log)
    with dest_engine.connect() as conn:
        assert conn.execute(text("select count(*) from bp_out")).scalar() == 8


def test_load_config_is_validated(project_copy):
    _edit(
        project_copy,
        "object.yaml",
        "loads:\n",
        "loads:\n"
        "  - kind: file\n    structure: ghost\n    path: ../../escape.csv\n    format: csv\n"
        "  - kind: table\n    structure: general\n    connection: nope\n    table: t\n",
    )
    with pytest.raises(ConfigError) as exc:
        load_project(project_copy)
    text_ = str(exc.value)
    assert "unknown structure 'ghost'" in text_
    assert "leaves the project" in text_
    assert "unknown connection 'nope'" in text_


def test_duplicate_load_destinations_are_rejected(project_copy):
    _edit(
        project_copy,
        "object.yaml",
        "loads:\n",
        "loads:\n"
        "  - kind: file\n    structure: general\n    path: output/supplier_general.csv\n"
        "    format: csv\n",
    )
    with pytest.raises(ConfigError, match="duplicate load destination"):
        load_project(project_copy)
