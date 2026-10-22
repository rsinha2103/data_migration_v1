"""Excel report."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FILLS = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "OK": PatternFill("solid", fgColor="C6EFCE"),
    "SUCCESS": PatternFill("solid", fgColor="C6EFCE"),
    "WARN": PatternFill("solid", fgColor="FFEB9C"),
    "ATTENTION": PatternFill("solid", fgColor="FFEB9C"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "FAILED": PatternFill("solid", fgColor="FFC7CE"),
    "ERROR": PatternFill("solid", fgColor="FFC7CE"),
    "BLOCKED": PatternFill("solid", fgColor="FFC7CE"),
}
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
MAX_CELL = 32_000


def _set(ws: Worksheet, row: int, col: int, value: Any) -> Any:
    """Write a cell. Text is always stored as text: values from migration data that start with
    '=' must never become formulas."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = value[:MAX_CELL]
        cell.data_type = "s"
    elif isinstance(value, bool | int | float) or value is None:
        cell.value = value
    else:
        cell.value = str(value)[:MAX_CELL]
        cell.data_type = "s"
    return cell


def _table(
    ws: Worksheet, headers: list[str], rows: list[list[Any]], status_col: str | None
) -> None:
    for c, h in enumerate(headers, start=1):
        cell = _set(ws, 1, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = _set(ws, r, c, value)
            if headers[c - 1] == status_col and value in FILLS:
                cell.fill = FILLS[value]
    ws.freeze_panes = "A2"
    for c, h in enumerate(headers, start=1):
        longest = max(
            [len(str(h))] + [len(str(r[c - 1])) for r in rows[:200] if r[c - 1] is not None]
        )
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 8), 60)


def _label(check: dict[str, Any]) -> str:
    if check["status"] == "FAIL" and check["severity"] == "warn":
        return "WARN"
    return str(check["status"])


def write_excel(report: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Summary"
    s = report["summary"]
    rows: list[tuple[str, Any]] = [
        ("Project", report["project"]),
        ("Object", report["object"]),
        ("Generated", report["generated_at"]),
        ("Tool", report["tool"]),
        ("Overall", report["overall"]),
        ("Checks passed", f"{s['checks_passed']} of {s['checks_total']}"),
        ("Checks blocking", s["checks_blocking"]),
        ("Check warnings", s["checks_warnings"]),
        ("Reconciliations passed", f"{s['recons_passed']} of {s['recons_total']}"),
        ("Reconciliations blocking", s["recons_blocking"]),
        ("Reconciliation warnings", s["recons_warnings"]),
        ("Failed steps", s["failed_steps"]),
        ("Runs included", ", ".join(r["run_id"] for r in report["runs"])),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        _set(ws, i, 1, k).font = Font(bold=True)
        cell = _set(ws, i, 2, v)
        if k == "Overall":
            cell.fill = FILLS.get(str(v), cell.fill)
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70
    ws.cell(row=1, column=2).alignment = Alignment(wrap_text=True)

    _table(
        wb.create_sheet("Steps"),
        ["Run", "Step", "Status", "Rows in", "Rows out", "Started", "Finished", "Message"],
        [
            [
                st["run_id"][:8],
                st["step_name"],
                st["status"],
                st["rows_in"],
                st["rows_out"],
                st["started_at"],
                st["finished_at"],
                st["message"],
            ]
            for st in report["steps"]
        ],
        "Status",
    )
    check_headers = [
        "Check", "Type", "Structure", "Field", "Severity", "Result", "Rows checked", "Failures",
        "Truncated", "Message",
    ]  # fmt: skip

    def check_rows(items: list[dict[str, Any]]) -> list[list[Any]]:
        return [
            [
                c["check_name"], c["check_type"], c["structure"], c["field"], c["severity"],
                _label(c), c["checked_count"], c["failed_count"], bool(c["truncated"]),
                c["message"],
            ]
            for c in items
        ]  # fmt: skip

    _table(wb.create_sheet("Checks"), check_headers, check_rows(report["checks"]), "Result")
    _table(
        wb.create_sheet("Reconciliation"),
        check_headers,
        check_rows(report["reconciliations"]),
        "Result",
    )
    _table(
        wb.create_sheet("Rejects"),
        ["Check", "Structure", "Severity", "Source row", "Field", "Value", "Message"],
        [
            [r["check_name"], r["structure"], r["severity"], r["row_no"], r["field"], r["value"],
             r["message"]]
            for r in report["rejects"]
        ],
        None,
    )  # fmt: skip
    profile_rows: list[list[Any]] = []
    for table, cols in report["profiles"].items():
        for p in cols:
            pct = 100.0 * p["empty_count"] / p["row_count"] if p["row_count"] else 0.0
            top = ", ".join(f"{v} ({n})" for v, n in json.loads(p["top_values"] or "[]")[:5])
            shapes = ", ".join(f"{v} ({n})" for v, n in json.loads(p["top_patterns"] or "[]")[:5])
            profile_rows.append(
                [table, p["column_name"], p["row_count"], p["empty_count"], round(pct, 1),
                 p["distinct_count"], p["min_value"], p["max_value"], p["min_length"],
                 p["max_length"], top, shapes]
            )  # fmt: skip
    _table(
        wb.create_sheet("Profile"),
        ["Table", "Column", "Rows", "Empty", "Empty %", "Distinct", "Min", "Max", "Min len",
         "Max len", "Top values", "Top shapes"],
        profile_rows,
        None,
    )  # fmt: skip
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
