"""Readers turning a source definition into chunks of polars DataFrames."""

from __future__ import annotations

from collections.abc import Iterator, Sequence
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import load_workbook
from sqlalchemy import Engine, literal_column, select, table, text

from data_migration_concept.config.models import FileSource, SqlSource
from data_migration_concept.extract.errors import ExtractError

CHUNK_SIZE = 10_000


def read_file(source: FileSource, object_dir: Path) -> Iterator[pl.DataFrame]:
    path = object_dir / source.path
    if not path.is_file():
        raise ExtractError(f"source '{source.name}': file '{source.path}' not found")
    try:
        if source.format == "csv":
            yield _read_csv(path, source)
        elif source.format == "xlsx":
            yield _read_xlsx(path, source)
        else:
            yield pl.read_parquet(path)
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError(f"source '{source.name}': cannot read '{source.path}': {exc}") from exc


def _read_csv(path: Path, source: FileSource) -> pl.DataFrame:
    # Decode ourselves so legacy encodings (cp1252, latin-1) work; stage everything as text.
    text_data = path.read_bytes().decode(source.encoding)
    return pl.read_csv(text_data.encode("utf-8"), separator=source.delimiter, infer_schema=False)


def _cell_to_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)


def _read_xlsx(path: Path, source: FileSource) -> pl.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if source.sheet is not None and source.sheet not in wb.sheetnames:
            raise ExtractError(f"source '{source.name}': sheet '{source.sheet}' not found")
        ws = wb[source.sheet] if source.sheet is not None else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ExtractError(f"source '{source.name}': sheet is empty")
        # Trailing blank header cells (common in Excel) are ignored; blanks inside are errors.
        header = list(header_row)
        while header and header[-1] is None:
            header.pop()
        if any(h is None or str(h).strip() == "" for h in header):
            raise ExtractError(f"source '{source.name}': a column in the header row has no name")
        names = [str(h).strip() for h in header]
        data: list[list[str | None]] = [[] for _ in names]
        for row in rows:
            cells = [_cell_to_text(c) for c in row[: len(names)]]
            if all(c is None for c in cells):
                continue
            cells += [None] * (len(names) - len(cells))
            for col, value in zip(data, cells, strict=True):
                col.append(value)
        return pl.DataFrame(
            {n: pl.Series(n, col, dtype=pl.Utf8) for n, col in zip(names, data, strict=True)}
        )
    finally:
        wb.close()


def read_sql(source: SqlSource, engine: Engine) -> Iterator[pl.DataFrame]:
    if source.table is not None:
        schema, _, name = source.table.rpartition(".")
        stmt: Any = select(literal_column("*")).select_from(table(name, schema=schema or None))
    else:
        stmt = text(source.query or "")
    try:
        with engine.connect().execution_options(stream_results=True) as conn:
            result = conn.execute(stmt)
            keys = list(result.keys())
            first = True
            while True:
                rows = result.fetchmany(CHUNK_SIZE)
                if not rows:
                    break
                first = False
                yield rows_to_frame(keys, rows)
            if first:  # no rows: still yield the column layout
                yield pl.DataFrame({k: pl.Series(k, [], dtype=pl.Utf8) for k in keys})
    except Exception as exc:
        raise ExtractError(f"source '{source.name}': query failed: {exc}") from exc


def rows_to_frame(keys: list[str], rows: Sequence[Any]) -> pl.DataFrame:
    columns = {k: [row[i] for row in rows] for i, k in enumerate(keys)}
    return pl.DataFrame(columns, infer_schema_length=None, strict=False)
