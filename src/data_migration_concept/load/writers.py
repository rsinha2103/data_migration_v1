"""Read a target table in chunks and write it to a file or a database table."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import polars as pl
import sqlalchemy as sa
from openpyxl import Workbook
from sqlalchemy import Engine, inspect, text

from data_migration_concept.config.models import FileLoad, TableLoad
from data_migration_concept.extract.readers import CHUNK_SIZE, rows_to_frame
from data_migration_concept.extract.staging import ROW_NO, RUN_ID, sa_type
from data_migration_concept.load.errors import LoadError
from data_migration_concept.transform.generator import q

LINEAGE = {ROW_NO, RUN_ID}


def table_columns(engine: Engine, table: str, include_lineage: bool) -> list[str]:
    names = [c["name"] for c in inspect(engine).get_columns(table)]
    return names if include_lineage else [n for n in names if n not in LINEAGE]


def row_count(engine: Engine, table: str) -> int:
    with engine.connect() as conn:
        return int(conn.execute(text(f"SELECT COUNT(*) FROM {q(table)}")).scalar_one())


def iter_frames(engine: Engine, table: str, columns: list[str]) -> Iterator[pl.DataFrame]:
    """Yield the table as DataFrames in source-row order; at least one (maybe empty) frame."""
    cols = ", ".join(q(c) for c in columns)
    sql = f"SELECT {cols} FROM {q(table)} ORDER BY {q(ROW_NO)}"
    yielded = False
    with engine.connect().execution_options(stream_results=True) as conn:
        result = conn.execute(text(sql))
        while True:
            rows = result.fetchmany(CHUNK_SIZE)
            if not rows:
                break
            yielded = True
            yield rows_to_frame(columns, rows)
    if not yielded:
        yield pl.DataFrame({c: pl.Series(c, [], dtype=pl.Utf8) for c in columns})


def write_file(frames: Iterator[pl.DataFrame], load: FileLoad, dest: Path) -> int:
    """Write frames to `dest` (the caller passes a temporary path). Returns the row count."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    rows = 0
    if load.format == "csv":
        with dest.open("wb") as fh:
            for i, frame in enumerate(frames):
                data = frame.write_csv(separator=load.delimiter, include_header=(i == 0))
                fh.write(data.encode(load.encoding))
                rows += frame.height
    elif load.format == "xlsx":
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(load.structure[:31])
        for i, frame in enumerate(frames):
            if i == 0:
                ws.append(frame.columns)
            for row in frame.iter_rows():
                ws.append(list(row))
            rows += frame.height
        wb.save(dest)
    else:
        collected = list(frames)
        pl.concat(collected, how="vertical_relaxed").write_parquet(dest)
        rows = sum(f.height for f in collected)
    return rows


def write_table(frames: Iterator[pl.DataFrame], load: TableLoad, engine: Engine) -> int:
    """Write frames to a database table in a single transaction. Returns the row count."""
    schema, _, name = load.table.rpartition(".")
    rows = 0
    tbl: sa.Table | None = None
    with engine.begin() as conn:
        for frame in frames:
            if tbl is None:
                tbl = sa.Table(
                    name,
                    sa.MetaData(),
                    *(sa.Column(n, sa_type(t), nullable=True) for n, t in frame.schema.items()),
                    schema=schema or None,
                )
                if load.mode == "replace":
                    tbl.drop(conn, checkfirst=True)
                    tbl.create(conn)
                else:
                    tbl.create(conn, checkfirst=True)
            if frame.height:
                conn.execute(sa.insert(tbl), frame.to_dicts())
                rows += frame.height
    if tbl is None:  # pragma: no cover - iter_frames always yields once
        raise LoadError(f"nothing to write to {load.table}")
    return rows
